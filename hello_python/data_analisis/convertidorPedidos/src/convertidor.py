from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

# Rutas del proyecto y de los archivos de trabajo.
BASE_DIR = Path(__file__).resolve().parent.parent
WORKSPACE_DIR = BASE_DIR.parent
RUTA_PRODUCTOS = r"C:\Users\Personal\Documents\script_filtro\productos.xlsx"
RUTA_PLANTILLA = Path(__file__).resolve().parent / "data" / "input" / "plantilla_pedidos.xlsx"
RUTA_SALIDA = BASE_DIR / "data" / "outputPedidos.xlsx"

COLUMNAS_REQUERIDAS = ["Nombre Cliente", "Nombre Producto", "Cantidad Pedida"]


class ConversionError(Exception):
    """Error controlado que se muestra al usuario."""


def procesar_pedidos(
    ruta_pedidos,
    ruta_productos=RUTA_PRODUCTOS,
    ruta_plantilla=RUTA_PLANTILLA,
    salida=RUTA_SALIDA,
):
    """Convierte el archivo de pedidos a la plantilla.

    Devuelve un dict con: clientes, salida, no_mapeados.
    Lanza ConversionError con un mensaje legible si algo falla.
    """
    ruta_pedidos = Path(ruta_pedidos)
    for ruta, nombre in (
        (ruta_pedidos, "pedidos"),
        (Path(ruta_productos), "productos (mapeo)"),
        (Path(ruta_plantilla), "plantilla"),
    ):
        if not ruta.exists():
            raise ConversionError(f"No se encontró el archivo de {nombre}: {ruta}")

    # 1. LEER EL ARCHIVO DE PEDIDOS
    # Saltamos las primeras 9 filas donde están los parámetros de búsqueda
    try:
        df = pd.read_excel(ruta_pedidos, skiprows=9)
    except Exception as e:
        raise ConversionError(f"No se pudo leer el archivo de pedidos: {e}") from e

    # Limpiamos nombres de columnas
    df.columns = df.columns.astype(str).str.strip()

    faltantes = [c for c in COLUMNAS_REQUERIDAS if c not in df.columns]
    if faltantes:
        raise ConversionError(
            "El archivo de pedidos no tiene las columnas: " + ", ".join(faltantes)
        )

    # Filtramos filas válidas
    df = df.dropna(subset=['Nombre Cliente', 'Nombre Producto'])
    if df.empty:
        raise ConversionError("El archivo de pedidos no contiene filas válidas.")

    # 2. PIVOT TABLE: Consolidar pedidos por cliente
    # Agrupamos por cliente y sumamos la cantidad pedida de cada producto
    pivot = df.pivot_table(
        index='Nombre Cliente',
        columns='Nombre Producto',
        values='Cantidad Pedida',
        aggfunc='sum',
        fill_value=0
    ).reset_index()

    # 3. DICCIONARIO DE MAPEO (Producto del sistema -> Nombre en la plantilla)
    df_map = pd.read_excel(ruta_productos)
    if not {'PRODUCTOS-SISTEMA', 'COLUMNA-PLANTILLA'}.issubset(df_map.columns):
        raise ConversionError(
            "El archivo de productos debe tener las columnas "
            "'PRODUCTOS-SISTEMA' y 'COLUMNA-PLANTILLA'."
        )
    mapeo_productos = dict(zip(df_map['PRODUCTOS-SISTEMA'],  df_map['COLUMNA-PLANTILLA']))

    # Extracción de los productos mapeados en el dict
    productos_en_pedido = set(df['Nombre Producto'].unique())

    productos_mapeados = set(mapeo_productos.keys())

    no_mapeados = sorted(str(p) for p in productos_en_pedido - productos_mapeados)

    if no_mapeados:
      print("\n" + "="*50)
      print("⚠️  ¡ATENCIÓN! PRODUCTOS NO REGISTRADOS EN EL MAPEO:")
      for prod in no_mapeados:
        print(f"❌ {prod}")
      print("Estos productos se omitirán en la plantilla hasta que los agregues.")


    # 4. ABRIR LA PLANTILLA CON OPENPYXL
    wb_plantilla = openpyxl.load_workbook(ruta_plantilla)
    ws: Worksheet = wb_plantilla.active  # type: ignore
    if ws is None:
        raise ConversionError("No se pudo cargar la hoja de la plantilla")

    # Crear un mapa de {Nombre de Columna: Número de Columna} de la plantilla
    col_map = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        if val:
            col_map[str(val).strip().lower()] = col

    # 5. ESCRIBIR EN LA PLANTILLA
    fila_inicio = 2  # Los datos inician en la fila 2

    for i, (_, row) in enumerate(pivot.iterrows()):
        fila_actual = fila_inicio + i
        cliente_nombre = str(row['Nombre Cliente'])

        # Escribir el nombre del cliente en la Columna 8 (H)
        ws.cell(row=fila_actual, column=8, value=cliente_nombre)

        # Opcional: Escribir Ruta en Columna 6 (F) si la tienes identificada
        # ws.cell(row=fila_actual, column=6, value="Ruta 1")

        # Escribir las cantidades de cada producto
        for prod_sistema, header_plantilla in mapeo_productos.items():
            if prod_sistema in row:
                cantidad = row[prod_sistema]
                if cantidad > 0:
                    # Buscar la columna en la plantilla
                    col_idx = col_map.get(str(header_plantilla).strip().lower())
                    if col_idx:
                        ws.cell(row=fila_actual, column=col_idx, value=cantidad)

    # 6. GUARDAR EL RESULTADO
    salida = Path(salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb_plantilla.save(salida)
    except PermissionError as e:
        raise ConversionError(
            f"No se pudo guardar '{salida.name}'. ¿Está abierto en Excel? Ciérralo e inténtalo de nuevo."
        ) from e

    print(f"¡Listo! Se procesaron {len(pivot)} clientes y se guardó en '{salida}' sin afectar fórmulas.")
    return {"clientes": len(pivot), "salida": str(salida), "no_mapeados": no_mapeados}


if __name__ == "__main__":
    procesar_pedidos(WORKSPACE_DIR / "pedidos.xlsx")
