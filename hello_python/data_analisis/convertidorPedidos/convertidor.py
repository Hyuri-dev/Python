import pandas as pd
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

# 1. LEER EL ARCHIVO DE PEDIDOS
# Saltamos las primeras 9 filas donde están los parámetros de búsqueda
df = pd.read_excel('pedidos.xlsx', skiprows=9)

# Limpiamos nombres de columnas
df.columns = df.columns.astype(str).str.strip()

# Filtramos filas válidas
df = df.dropna(subset=['Nombre Cliente', 'Nombre Producto'])

# 2. PIVOT TABLE: Consolidar pedidos por cliente
# Agrupamos por cliente y sumamos la cantidad pedida de cada producto
pivot = df.pivot_table(
    index='Nombre Cliente',
    columns='Nombre Producto',
    values='Cantidad Pedida',
    aggfunc='sum',
    fill_value=0
).reset_index()

# 3. DICCIONARIO DE MAPEO (Producto del sistema -> Nombre en la plantilla)
ruta_productos = r"C:\Users\Personal\Documents\script_filtro\productos.xlsx"
df_map = pd.read_excel(ruta_productos)
mapeo_productos = dict(zip(df_map['PRODUCTOS-SISTEMA'],  df_map['COLUMNA-PLANTILLA']))

# Extracción de los productos mapeados en el dict
productos_en_pedido = set(df['Nombre Producto'].unique())

productos_mapeados = set(mapeo_productos.keys())

no_mapeados = productos_en_pedido - productos_mapeados

if no_mapeados:
  print("\n" + "="*50)
  print("⚠️  ¡ATENCIÓN! PRODUCTOS NO REGISTRADOS EN EL MAPEO:")
  for prod in no_mapeados:
    print(f"❌ {prod}")
  print("Estos productos se omitirán en la plantilla hasta que los agregues.")

  
# 4. ABRIR LA PLANTILLA CON OPENPYXL
wb_plantilla = openpyxl.load_workbook('plantilla_pedidos.xlsx')
ws: Worksheet = wb_plantilla.active  # type: ignore
assert ws is not None, "No se pudo cargar la hoja de la plantilla"

# Crear un mapa de {Nombre de Columna: Número de Columna} de la plantilla
col_map = {}
for col in range(1, ws.max_column + 1):
    val = ws.cell(row=1, column=col).value
    if val:
        col_map[str(val).strip().lower()] = col

# 5. ESCRIBIR EN LA PLANTILLA
fila_inicio = 2  # Los datos inician en la fila 2

for i, (_, row) in enumerate(pivot.iterrows()):
    fila_actual = fila_inicio + i
    cliente_nombre = str(row['Nombre Cliente'])

    # Escribir el nombre del cliente en la Columna 8 (H)
    ws.cell(row=fila_actual, column=8, value=cliente_nombre)

    # Opcional: Escribir Ruta en Columna 6 (F) si la tienes identificada
    # ws.cell(row=fila_actual, column=6, value="Ruta 1")

    # Escribir las cantidades de cada producto
    for prod_sistema, header_plantilla in mapeo_productos.items():
        if prod_sistema in row:
            cantidad = row[prod_sistema]
            if cantidad > 0:
                # Buscar la columna en la plantilla
                col_idx = col_map.get(header_plantilla.strip().lower())
                if col_idx:
                    ws.cell(row=fila_actual, column=col_idx, value=cantidad)

# 6. GUARDAR EL RESULTADO
salida = 'plantilla_con_pedidos_cargados.xlsx'
wb_plantilla.save(salida)
print(f"¡Listo! Se procesaron {len(pivot)} clientes y se guardó en '{salida}' sin afectar fórmulas.")