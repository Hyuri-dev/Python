import pandas as pd
import openpyxl as opyxl
from openpyxl import load_workbook
from rich.console import Console
from rich.table import Table

# ----------- Headers del dataframe nuevo -----------
nombres_limpios = [
    "Codigo", "Descripcion", "Cantidad", "Pendiente", "MontoBruto", 
    "Descuentos", "IVA", "Costo", "Utilidad", "PorcUtilidad", "Existencia"
]

ubicacion_reporte = {
    "Julio Medina": r"\\SERVIDOR\a2Apps\a2Admin\Empre001\REPORTS\ProductosvendidosJULIO.TXT",
    "Victor Ferreira": r"\\SERVIDOR\a2Apps\a2Admin\Empre001\REPORTS\Productosvendidosvictorf.TXT",
    "Luis Duran": r"\\SERVIDOR\a2Apps\a2Admin\Empre001\REPORTS\ProductosvendidosLUIS.TXT",
    "Robert Rodrigues": r"\\SERVIDOR\a2Apps\a2Admin\Empre001\REPORTS\ProductosvendidosRobert.TXT",
    "Distribuidora": r"\\SERVIDOR\a2Apps\a2Admin\Empre001\REPORTS\ProductosvendidosDist.TXT" ,
    "Yosemith Ponce": r"\\SERVIDOR\a2Apps\a2Admin\Empre001\REPORTS\ProductosvendidosYOSE.TXT",
    "Jesus Hernandez": r"\\SERVIDOR\a2Apps\a2Admin\Empre001\REPORTS\ProductosvendidosJesus.TXT",
    "Victor Gonzales": r"\\SERVIDOR\a2Apps\a2Admin\Empre001\REPORTS\ProductosvendidosvictorG.TXT"
}

productos = {
    "Pasta Allegri": "PASTA ALLEGRI",
    "Pasta Horizonte": "HORIZONTE",
    "Pasticho Allegri": "PASTICHO ALLEGRI",
    "Pasticho Mi Casa": "PASTICHO MI CASA",
    "Harina de Trigo Dulce Mar":  "HARINA DE TRIGO DULCE MAR",
    "Harina de Maiz Juana": "HARINA JUANA",
    "Harina de Trigo Robin Hood":"ROBIN HOOD",
    "Harina de Cachapa Juana": "MEZCLA CACHAPA JUANA",
    "Arroz Monica": "ARROZ MONICA",
    "Chococao": "BEBIDA CHOCOCAO",
    "Margarina Juana":"MARGARINA JUANA",
    "Aceite Vegetal": "ACEITE VEGETAL LA COMADRE",
    "Avena Robin Hood":"AVENA ROBIN HOOD HOJ BOLSA"
}

vendedores_mapeo = {
    "Julio Medina": 6,
    "Victor Ferreira": 7,
    "Luis Duran": 8,
    "Robert Rodrigues": 9,
    "Distribuidora": 10 ,
    "Yosemith Ponce": 11,
    "Jesus Hernandez": 12,
    "Victor Gonzales": 13
}

ubicacion_productos = {
    "Pasta Allegri": "I",
    "Pasta Horizonte": "G",
    "Harina de Trigo Dulce Mar": "K",
    "Pasticho Allegri": "M",
    "Pasticho Mi Casa": "O"
}

productos_por_codigo = ['001009', '001011' , '001012']
allegri = ["Pasta Allegri", "Pasta Horizonte", "Pasticho Allegri","Allegri Especialidades" ,"Pasticho Mi Casa", "Harina de Trigo Dulce Mar"]
monaca = ["Harina de Maiz Juana","Harina de Cachapa Juana" ,"Harina de Trigo Robin Hood", "Arroz Monica", "Chococao", "Margarina Juana", "Aceite Vegetal"]

def crea_reporte():
    reporte_resumen = []
    ruta_excel = r"C:\Users\Personal\Documents\JEFF\NOVIEMBRE 2025\DIACENCA_VENTAS_PRUEBA.xlsm"
    
    try:
        # Abrimos el Excel antes de empezar el bucle para escribir poco a poco
        libro = opyxl.load_workbook(ruta_excel, keep_vba=True)
        ws = libro["Hoja1"]
        
        for vendedor, ruta in ubicacion_reporte.items():
            print(f"Procesando a: {vendedor}")
            df = pd.read_csv(
                ruta, sep='\t', encoding='latin-1', names=nombres_limpios,
                header=0, thousands='.', decimal=',', index_col=False, 
                dtype={'Codigo': str}
            )
            
            # Identificar fila del vendedor
            fila_vendedor = vendedores_mapeo.get(vendedor)

            for nombre_producto, texto_buscar in productos.items():
                filtro = df[(df['Descripcion'].str.contains(texto_buscar, case=False, na=False)) & (df['Cantidad'] > 0)]
                total_cantidad = filtro['Cantidad'].sum()
                total_bruto = filtro['MontoBruto'].sum()
                IVA = filtro['IVA'].sum()

                reporte_resumen.append({
                    "Vendedor": vendedor,
                    "Producto/Categoria": nombre_producto,    
                    "Cantidad Total": total_cantidad,
                    "MontoBruto": total_bruto,
                    "IVA": IVA
                })

                # --- ESCRITURA EN EXCEL ---
                # Si el producto tiene columna asignada y el vendedor tiene fila
                columna_letra = ubicacion_productos.get(nombre_producto)
                if fila_vendedor and columna_letra:
                    ws[f"{columna_letra}{fila_vendedor}"] = total_cantidad

        # Guardar cambios al terminar todos los vendedores
        libro.save(ruta_excel)
        print(" ✅ Excel Guardado exitosamente")

    except Exception as e:
        print(f" ❌ Error en el proceso: {e}")

    # --- Lógica de Consola (Basada en tu DF final) ---
    df_reporte = pd.DataFrame(reporte_resumen)
    filtro_total = df_reporte['Cantidad Total'].sum()
    filtro_monto_global = df_reporte['MontoBruto'].sum()
    filtro_monto_allegri = df_reporte[df_reporte['Producto/Categoria'].isin(allegri)]['MontoBruto'].sum()
    filtro_monto_monaca = df_reporte[df_reporte['Producto/Categoria'].isin(monaca)]['MontoBruto'].sum()

    console = Console()
    table = Table(title="📊 Reporte De Ventas")
    table.add_column("Producto", style="cyan")
    table.add_column("Cant. Total", justify="right", style="magenta")
    table.add_column("Monto Bruto", justify="right", style="yellow")

    for _, row in df_reporte.iterrows():
        table.add_row(str(row['Producto/Categoria']), f"{row['Cantidad Total']:,.2f}", f"{row['MontoBruto']:,.2f}")
    
    table.add_row("TOTAL MONTO ALLEGRI", f"{filtro_monto_allegri:,.2f}", "-", style="bold yellow on green")
    table.add_row("TOTAL MONTO MONACA", f"{filtro_monto_monaca:,.2f}", "-", style="bold yellow on green")
    
    console.print(table)

crea_reporte()